from __future__ import annotations

from openpyxl import Workbook

from app.config import get_settings
from app.services.bid_service import get_autovad_master_template_lines


def test_master_template_loader_reads_offset_header_sheet(tmp_path, monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bid Items"
    ws["A1"] = "City of Sioux Falls Bid Item List"
    ws["A2"] = "Updated:"
    ws["A7"] = "Bid Item Number"
    ws["B7"] = "Unit"
    ws["C7"] = "Item Description"
    ws["A8"] = "4.001"
    ws["B8"] = "Hour"
    ws["C8"] = "Blading"
    ws["A9"] = "4.002"
    ws["B9"] = "LS"
    ws["C9"] = "Construction and Maintenance of Detour(s)"
    path = tmp_path / "Bid Item List 2026.xlsx"
    wb.save(path)
    wb.close()

    monkeypatch.setenv("AUTOVAD_MASTER_BID_TEMPLATE_PATH", str(path))
    monkeypatch.setenv("AUTOVAD_MASTER_BID_TEMPLATE_SHEET", "Bid Items")
    get_settings.cache_clear()
    try:
        name, lines, err = get_autovad_master_template_lines(force_reload=True)
        assert err is None
        assert "bid item list 2026.xlsx" in name.lower()
        assert len(lines) == 2
        assert lines[0].item_code == "4.001"
        assert lines[0].description == "Blading"
        assert lines[0].unit.lower() == "hour"
        assert lines[1].item_code == "4.002"
        assert lines[1].unit.lower() == "ls"
    finally:
        get_settings.cache_clear()
