from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from app.services.eoq_eval import compare_eoq
from app.services.eoq_service import EOQ_EXPORT_HEADERS
from app.services.training_service import parse_autovad_eoq_file


def _write_autovad_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate Of Quantities"
    ws["A1"] = "ESTIMATE OF QUANTITIES"
    for col, header in enumerate(EOQ_EXPORT_HEADERS, start=1):
        ws.cell(2, col, header)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(EOQ_EXPORT_HEADERS))
    ws.cell(3, 1, "Watermain")
    ws.cell(4, 1, 1)
    ws.cell(4, 2, "8.110")
    ws.cell(4, 3, "8-Inch PVC Water Main")
    ws.cell(4, 4, "LF")
    ws.cell(4, 5, 240)
    ws.cell(4, 8, 98.5)
    ws.cell(4, 11, "User portal export")

    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(EOQ_EXPORT_HEADERS))
    ws.cell(5, 1, "Traffic")
    ws.cell(6, 1, 2)
    ws.cell(6, 2, "Special")
    ws.cell(6, 3, "Traffic Control")
    ws.cell(6, 4, "SQFT")
    ws.cell(6, 5, 624)
    ws.cell(6, 8, 91.2)
    wb.save(path)
    wb.close()


def test_parse_autovad_excel_reads_export_layout(tmp_path):
    path = tmp_path / "autovad_eoq.xlsx"
    _write_autovad_workbook(path)
    items = parse_autovad_eoq_file(path, path.name)
    assert len(items) == 2
    assert items[0]["description"] == "8-Inch PVC Water Main"
    assert items[0]["unit"] == "LF"
    assert items[0]["quantity"] == 240
    assert items[0]["category"] == "Watermain"
    assert items[0]["item_code"] == "8.110"
    assert items[0]["confidence"] == 98.5
    assert items[1]["description"] == "Traffic Control"
    assert items[1]["category"] == "Traffic"
    assert items[1]["item_code"] == "Special"


def test_parse_autovad_csv_reads_group_banners(tmp_path):
    path = tmp_path / "autovad_eoq.csv"
    path.write_text(
        "\ufeffItem Number,Group,Standard Bid Item Number,Item Description,Unit,Quantity,AI Confidence\n"
        ",Watermain,,===== Watermain =====,,,\n"
        "1,Watermain,8.110,8-Inch PVC Water Main,LF,240,98.50\n"
        ",Traffic,,===== Traffic =====,,,\n"
        "2,Traffic,Special,Traffic Control,SQFT,624,91.20\n",
        encoding="utf-8",
    )
    items = parse_autovad_eoq_file(path, path.name)
    assert [i["description"] for i in items] == ["8-Inch PVC Water Main", "Traffic Control"]
    assert items[0]["category"] == "Watermain"
    assert items[1]["unit"] == "SQFT"
    assert items[1]["quantity"] == 624


def test_parse_autovad_simple_excel_headers(tmp_path):
    path = tmp_path / "simple.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Description", "Unit", "Quantity", "Category"])
    ws.append(["Remove Concrete Pavement", "SY", 120, "Removals"])
    wb.save(path)
    wb.close()
    items = parse_autovad_eoq_file(path, path.name)
    assert len(items) == 1
    assert items[0]["description"] == "Remove Concrete Pavement"
    assert items[0]["unit"] == "SY"
    assert items[0]["quantity"] == 120
    assert items[0]["category"] == "Removals"


def test_imported_autovad_items_compare_against_original():
    actual = [
        {
            "description": "8-Inch PVC Water Main",
            "unit": "LF",
            "quantity": 240,
            "category": "Watermain",
            "item_code": "8.110",
        }
    ]
    expected = {
        "items": [
            {
                "description": "8-Inch PVC Water Main",
                "unit": "LF",
                "quantity": 240,
                "category": "Watermain",
            }
        ]
    }
    report = compare_eoq(expected, actual)
    assert report.expected_count == 1
    assert report.actual_count == 1
    assert not report.misses
    assert not report.extras


def test_parse_autovad_rejects_empty_workbook(tmp_path):
    path = tmp_path / "empty.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ESTIMATE OF QUANTITIES"
    wb.save(path)
    wb.close()
    assert parse_autovad_eoq_file(path, path.name) == []
