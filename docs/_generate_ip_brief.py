"""Generate AutoVAD patent-style unique-engines pack with red CONFIDENTIAL watermark."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = Path(__file__).resolve().parent / "AutoVAD_Takeoff_Engines_Confidential_IP_Brief.docx"

ENGINE_FILES = [
    "backend/app/services/ai_analysis.py",
    "backend/app/services/incidental.py",
    "backend/app/services/item_combine.py",
    "backend/app/services/eoq_validation.py",
    "backend/app/services/traffic_control.py",
    "backend/app/data/mutcd_sign_sizes.json",
    "backend/app/services/civil_estimator.py",
    "backend/app/services/utility_labels.py",
    "backend/app/services/csi_mapper.py",
    "backend/app/services/eoq_groups.py",
    "backend/app/services/eoq_service.py",
    "backend/app/services/bid_service.py",
    "backend/app/services/eoq_eval.py",
    "backend/app/services/pdf_vision.py",
    "backend/app/services/openai_client.py",
    "backend/app/services/extractors.py",
    "backend/app/services/cad/quantity_engine.py",
    "backend/app/services/cad/utility_stationing.py",
    "backend/app/services/cad/civil_location.py",
    "backend/app/services/cad/engine.py",
    "backend/app/services/cad/dxf_parser.py",
    "backend/app/services/cad/landxml_parser.py",
    "backend/app/services/cad/aps_client.py",
    "backend/app/services/cad/design_automation.py",
    "backend/app/services/cad/dwg_aps.py",
    "backend/app/services/cad/autodesk_aps.py",
    "backend/cad_plugins/AutoVadCivilTakeoff/Commands.cs",
    "backend/cad_plugins/AutoVadCivilTakeoff/PackageContents.xml",
    "frontend/src/utils/eoqGroups.ts",
    "backend/app/schemas/units.py",
]

RED = RGBColor(192, 0, 0)
DARK = RGBColor(13, 31, 25)
BLUE = RGBColor(31, 78, 121)

HDR_XML = """<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
       xmlns:v="urn:schemas-microsoft-com:vml"
       xmlns:o="urn:schemas-microsoft-com:office:office"
       xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <w:p>
    <w:pPr>
      <w:pStyle w:val="Header"/>
      <w:jc w:val="center"/>
    </w:pPr>
    <w:r>
      <w:rPr>
        <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>
        <w:b/>
        <w:color w:val="C00000"/>
        <w:sz w:val="20"/>
      </w:rPr>
      <w:t>CONFIDENTIAL</w:t>
    </w:r>
  </w:p>
  <w:p>
    <w:r>
      <w:rPr><w:noProof/></w:rPr>
      <w:pict>
        <v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800"
          path="m@7,l@8,m@5,21600l@6,21600e">
          <v:formulas>
            <v:f eqn="sum #0 0 10800"/>
            <v:f eqn="prod #0 2 1"/>
            <v:f eqn="sum 21600 0 @1"/>
            <v:f eqn="sum 0 0 @2"/>
            <v:f eqn="sum 21600 0 @3"/>
            <v:f eqn="if @0 @3 0"/>
            <v:f eqn="if @0 21600 @1"/>
            <v:f eqn="if @0 0 @2"/>
            <v:f eqn="if @0 @4 21600"/>
            <v:f eqn="mid @5 @6"/>
            <v:f eqn="mid @8 @5"/>
            <v:f eqn="mid @7 @8"/>
            <v:f eqn="mid @6 @7"/>
            <v:f eqn="sum @6 0 @5"/>
          </v:formulas>
          <v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800"/>
          <v:textpath on="t" fitshape="t"/>
          <v:handles>
            <v:h position="#0,bottomRight" xrange="6629,14971"/>
          </v:handles>
        </v:shapetype>
        <v:shape id="AutoVADConfidentialWM" o:spid="_x0000_s2049" type="#_x0000_t136"
          alt="CONFIDENTIAL"
          style="position:absolute;margin-left:0;margin-top:0;width:500pt;height:140pt;rotation:315;z-index:-251658752;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
          o:allowincell="f" filled="t" fillcolor="red" stroked="f">
          <v:fill opacity=".38"/>
          <v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="CONFIDENTIAL"/>
        </v:shape>
      </w:pict>
    </w:r>
  </w:p>
</w:hdr>"""


def set_run_font(run, *, size=11, bold=False, color=None, italic=False, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def apply_header(section):
    for header in (section.header, section.even_page_header):
        header.is_linked_to_previous = False
        parent = header._element
        parent.clear()
        hdr = parse_xml(HDR_XML)
        for child in list(hdr):
            parent.append(child)


def apply_footer(section):
    for footer in (section.footer, section.even_page_footer):
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def add_red(text: str):
            run = p.add_run(text)
            set_run_font(run, size=9, bold=True, color=RED)
            return run

        add_red("CONFIDENTIAL  ·  AutoVAD unique source  ·  Page ")
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        r = p.add_run()
        r._r.append(fld1)
        r._r.append(instr)
        r._r.append(fld2)
        set_run_font(r, size=9, bold=True, color=RED)
        add_red("  ·  Do not copy or distribute")


def centered(doc, text, *, size=11, bold=False, color=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def para(doc, text, *, italic=False, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run_font(run, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(cell, "0D1F19")
    for r_i, row in enumerate(rows):
        bg = "F4F7F2" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=9, name="Consolas" if c_i > 0 else "Calibri")
            shade_cell(cell, bg)
    doc.add_paragraph()
    return table


def file_stats(rel: str) -> tuple[int, str, str]:
    path = ROOT / rel
    data = path.read_bytes()
    text = data.decode("utf-8", errors="replace")
    lines = text.count("\n") + (0 if text.endswith("\n") or not text else 1)
    digest = hashlib.sha256(data).hexdigest()
    return lines, digest, text


def add_source(doc, rel: str, lines: int, digest: str, text: str):
    doc.add_heading(rel, level=3)
    para(
        doc,
        f"SHA-256: {digest}    ·    {lines} line(s)    ·    {Path(rel).suffix or 'file'}",
        size=9,
        space_after=4,
    )
    chunk_size = 3500
    for start in range(0, len(text), chunk_size):
        chunk = text[start : start + chunk_size]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(chunk)
        set_run_font(run, size=7, name="Consolas")


def build() -> Path:
    stats = []
    missing = []
    for rel in ENGINE_FILES:
        path = ROOT / rel
        if not path.exists():
            missing.append(rel)
            continue
        lines, digest, text = file_stats(rel)
        stats.append((rel, lines, digest, text))
    if missing:
        raise FileNotFoundError("Missing engine files:\n" + "\n".join(missing))

    total_files = len(stats)
    total_lines = sum(s[1] for s in stats)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.05)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.different_odd_and_even_pages = True
    apply_header(section)
    apply_footer(section)

    centered(doc, "CONFIDENTIAL", size=14, bold=True, color=RED)
    centered(doc, "ATTORNEY WORK PRODUCT  ·  PATENT REVIEW (UNIQUE ENGINES)", size=11, bold=True, color=RED)
    centered(doc, "AutoVAD", size=28, bold=True, color=DARK)
    centered(doc, "AI Copilot for Civil Engineers", size=16, color=BLUE)
    centered(doc, "Patent-Only Pack\nUnique Takeoff, CAD, and EOQ Engines", size=14, italic=True)
    centered(
        doc,
        f"Prepared: {date.today().isoformat()}    ·    Engine source files: {total_files}    ·    Approximate source lines: {total_lines:,}",
        size=10,
        color=RGBColor(80, 80, 80),
    )
    para(
        doc,
        "This pack is for patent counsel. It describes the novel automated civil-quantity methods and reprints "
        "only the engine source that implements them. It does not include the website shell (login, projects UI, "
        "billing, admin, notifications). This document is not a patent application and is not legal advice.",
    )

    doc.add_heading("1. Confidentiality and handling", level=1)
    para(
        doc,
        "Every page of this document is marked CONFIDENTIAL in red in the header, footer, and as a diagonal watermark. "
        "Counsel, inventors, and authorized reviewers only. Do not email to personal accounts, upload to public AI tools, "
        "or share with vendors without a written NDA.",
    )
    bullets(
        doc,
        [
            "Classification: Confidential — trade secret / unpublished source.",
            "Suggested legend: CONFIDENTIAL — AutoVAD — Attorney review only.",
            "Use this pack for patent claims. Use a full source deposit for copyright registration of the whole program.",
        ],
    )

    doc.add_heading("2. What this pack is for", level=1)
    para(
        doc,
        "Patent law protects novel methods and systems, not every line of a website. This pack is limited to AutoVAD’s "
        "unique quantity-takeoff engines. Owner / inventors: to be completed by counsel.",
    )
    bullets(
        doc,
        [
            "Included: PDF/CAD takeoff engines, incidental filter, location combining, traffic-control rollup, CSI/EOQ grouping including Alternate A/B, master bid-template matching, deterministic validation, Civil 3D stationing, Excel EOQ engine.",
            "Excluded: login, registration, dashboard chrome, billing, notifications, admin UI, generic API glue.",
        ],
    )

    doc.add_heading("3. System overview (engine)", level=1)
    para(
        doc,
        "An engineer uploads PDF plan sets and/or CAD (DXF, DWG, LandXML). The engines extract pay items and quantities, "
        "build an Estimate of Quantities (EOQ), map evidenced items to AutoVAD’s master bid template, and export Excel. "
        "Schedule tables are treated as authoritative when present. Unmatched evidenced takeoff is kept with Standard Bid "
        "Item Number shown as Special, still grouped under the correct civil category. Alternate A / Alternate B sections "
        "are separate EOQ categories.",
    )
    add_table(
        doc,
        ["Engine", "Role"],
        [
            ["Document AI fusion", "Text, tables, PDF vision, labels — schedule-first ranking and optional strict schedule lock"],
            ["Incidental filter", "Drop work marked incidental; do not add it into the parent qty"],
            ["Location combiner", "Sum similar pay items from different sheets; keep true variants and alternates apart"],
            ["Validation layer", "Unit-family sanity, duplicate collapse, schedule-priority over plan takeoff"],
            ["Traffic Control", "MUTCD sign rollup to SqFt, or copy schedule companion rows"],
            ["CAD / Civil 3D", "Size-aware quantities, trench extras, APS DWG plugin, station-offset lengths"],
            ["EOQ assembly", "CSI + agency bid numbers or Special, municipal groups, Alternate A/B, Mobilization, Excel review"],
        ],
    )

    doc.add_heading("4. Unique methods for patent review", level=1)
    para(
        doc,
        "Counsel should assess patentability (novelty, non-obviousness, eligible subject matter) and whether to file a "
        "provisional covering the combination of schedule-first extraction, incidental exclusion, location combining, "
        "traffic-control rollup, CAD stationing, master-template evidence matching, and municipal EOQ presentation.",
    )

    methods = [
        (
            "U-01  Bid-schedule-authoritative pay-item copy",
            "When an Estimate of Quantities / Bid Items table exists, AutoVAD copies those rows as the pay-item list, "
            "preserves blank unit/quantity cells, and drops F-sheet device tables, graphic symbol counts, assumed trench extras, "
            "and typical-section invents. Strict schedule lock is used only when rows look like a true bid schedule (item numbers / "
            "agency codes), not a general detail quantity table. Agency bid numbers (for example 9.0010 or 634.0110) are copied as item_code.",
            "backend/app/services/ai_analysis.py — _is_bid_schedule_table, _is_strict_schedule_lock_row, _items_from_document_tables, _finalize_analysis",
        ),
        (
            "U-02  Multi-engine drawing takeoff fusion",
            "PDF plans are read by combining OpenAI text/table extraction, rendered-sheet vision, deterministic utility-label parsing, "
            "heuristic patterns, and CAD geometry. Pages are scored so bid/qty/utility sheets are preferred; large sets can scan all pages "
            "in RAM-safe batches. Evidence is ranked so a printed schedule quantity is preferred over a drawing measurement.",
            "backend/app/services/ai_analysis.py, pdf_vision.py, openai_client.py, utility_labels.py",
        ),
        (
            "U-03  Incidental-to-bid-item exclusion",
            "Drawing notes that mark work as incidental to / included in / paid under a bid item (bedding, trench, fittings, tracer wire, "
            "thrust blocks, testing, and similar) are omitted as separate pay items and are not added into the parent quantity.",
            "backend/app/services/incidental.py and wiring in ai_analysis.py, utility_labels.py, civil_estimator.py",
        ),
        (
            "U-04  Similar-item location combining",
            "The same pay item taken off on multiple sheets (example: Fertilizer 1,189 lb and Fertilizer 39 lb) is identified as one item "
            "and the quantities are added. True variants stay separate (pipe size, barricade type, fertilizer N-P-K grade, seed mix, "
            "remove vs install, Alternate A vs Alternate B, and incompatible units).",
            "backend/app/services/item_combine.py — combine_similar_pay_items",
        ),
        (
            "U-05  Traffic Control sign rollup with schedule companions",
            "Individual MUTCD / plan signs are rolled into one Traffic Control SqFt pay item when no bid schedule exists. When a Bid Items "
            "table has a Traffic Control section, every printed companion row is copied (SqFt signing, miscellaneous LS, barricades) and "
            "plan-invented signs are dropped.",
            "backend/app/services/traffic_control.py; MUTCD sizes in backend/app/data/mutcd_sign_sizes.json",
        ),
        (
            "U-06  Size-aware CAD quantity engine",
            "CAD entities are converted to civil pay items by network (water / sanitary / storm), detected pipe size, length, hatch area, "
            "and named-block counts. Trench excavation / bedding / backfill CY may be derived from pipe OD with stated cover assumptions. "
            "Unit scaling to feet is recorded when the drawing is unitless.",
            "backend/app/services/cad/quantity_engine.py, civil_estimator.py",
        ),
        (
            "U-07  Civil 3D / DWG station-offset takeoff",
            "Native DWG is processed through Autodesk Platform Services and an AutoCAD Design Automation plugin (AutoVadCivilTakeoff) that "
            "exports alignments, pipes with start/end XY, blocks, and paper-space callouts. Station-offset logic builds centerline-relative "
            "LT/RT lengths, fittings, bid summary, and QA flags for Excel detail sheets.",
            "backend/app/services/cad/utility_stationing.py, aps_client.py, design_automation.py, dwg_aps.py; backend/cad_plugins/AutoVadCivilTakeoff/Commands.cs",
        ),
        (
            "U-08  Municipal EOQ grouping, alternates, and Mobilization",
            "Pay items are grouped in agency EOQ order (General, Traffic Control, Removals, Grading, utilities, surfacing, and related). "
            "Alternate A / Alternate B / Option sections are their own categories. Mobilization (1 LS) is stored once under General unless "
            "it belongs to an alternate. Alternate labels are stripped from item descriptions because the section header already identifies them.",
            "backend/app/services/eoq_groups.py; frontend/src/utils/eoqGroups.ts; eoq_service.ensure_mobilization_item",
        ),
        (
            "U-09  CSI mapping that does not erase agency bid numbers",
            "Descriptions are mapped to CSI MasterFormat codes and USA pay units (including SQFT and TON export labels), while agency "
            "Standard Bid Item Numbers remain on item_code when they look like bid numbers rather than CSI.",
            "backend/app/services/csi_mapper.py",
        ),
        (
            "U-10  Design-only civil estimator",
            "When no bid schedule exists, typical-section dimensions produce pavement CY/TON, prime/tack SY, building/dam quantities, and "
            "related estimator lines, with assumptions written into calculation_method for engineer review.",
            "backend/app/services/civil_estimator.py",
        ),
        (
            "U-11  Plan-label utility parser",
            "Callouts such as 8\" WATER MAIN 245 LF or station-range labels are parsed deterministically into sized water / sanitary / storm "
            "quantities. Fitting invents are suppressed when an EOQ schedule is present.",
            "backend/app/services/utility_labels.py",
        ),
        (
            "U-12  Master bid-template matching without dumping the unused list",
            "AutoVAD’s standard master bid list is matched to evidenced takeoff only. Unused template lines are omitted. Multiple location "
            "hits on the same template line are summed. Unmatched evidenced items stay in the EOQ with Standard Bid Item Number = Special, "
            "and are still placed under Watermain, Surfacing, or other true categories rather than a leftover unmapped bucket.",
            "backend/app/services/bid_service.py — build_eoq_items_from_template, _match_line",
        ),
        (
            "U-13  Deterministic EOQ validation",
            "After extraction and before persistence, AutoVAD enforces unit-family sanity (linear/area/volume/count/lump), collapses exact "
            "duplicates, and prefers schedule quantities over conflicting plan-derived rows of the same pay item.",
            "backend/app/services/eoq_validation.py — validate_extracted_items",
        ),
        (
            "U-14  Gold-set EOQ evaluation / Training Lab",
            "Expected EOQ gold cases are compared to engine output (recall, misses by category, quantity error). An internal Training Lab "
            "lets reviewers compare original vs AutoVAD analysis vs evaluation. Gold cases are used to improve engines; they are not hardcoded "
            "as the product EOQ for other projects.",
            "backend/app/services/eoq_eval.py, training_service.py; frontend training views",
        ),
        (
            "U-15  Engineer-review Excel EOQ workbook",
            "Excel export writes municipal section headers, bid numbers or Special, SQFT/TON labels, and conditional formatting so Verified vs "
            "Engineer Review changes cell fill and font. Utility stationing can export detail plus rolled summary.",
            "backend/app/services/eoq_service.py",
        ),
        (
            "U-16  Plan-page vision scoring for large sets",
            "Each PDF page is scored for civil takeoff value (schedule, utility labels, profiles, sparse-text drawings). Utility/schedule pages "
            "can be force-included. Large documents use smaller vision batches without silently dropping remaining sheets when scan-all is on.",
            "backend/app/services/pdf_vision.py — plan_pdf_vision_pages, _score_page",
        ),
    ]
    for title, body, source in methods:
        doc.add_heading(title, level=2)
        para(doc, body)
        para(doc, f"Primary source: {source}", italic=True, size=10)

    doc.add_heading("5. What is not claimed as AutoVAD invention", level=1)
    bullets(
        doc,
        [
            "Generic civil formulas (CY = L × W × D / 27) as such.",
            "MUTCD sign designations as a public standard; AutoVAD’s use/rollup logic is claimed, not the MUTCD catalog.",
            "CSI MasterFormat numbering as a published classification; AutoVAD’s mapping rules and agency-number preservation are claimed.",
            "Autodesk APS / AutoCAD / Civil 3D platforms; AutoVAD’s plugin, property mapping, and stationing pipeline are claimed.",
            "OpenAI models; AutoVAD’s prompts, fusion policy, and post-filters are claimed.",
            "Agency-published bid item lists as data; AutoVAD’s matching engine and evidence-only EOQ build are claimed.",
            "Ordinary website features (login, project CRUD, billing pages) — copyright only, not patent claims in this pack.",
        ],
    )

    doc.add_heading("6. Engine source inventory", level=1)
    para(
        doc,
        "Each engine file below is identified by path, line count, and SHA-256. Section 7 reprints those files in full.",
    )
    para(doc, f"Total engine files: {total_files}. Total lines: {total_lines:,}.")
    add_table(
        doc,
        ["Path", "Lines", "SHA-256 (prefix)"],
        [[rel, str(lines), digest[:16] + "…"] for rel, lines, digest, _text in stats],
    )

    doc.add_heading("7. Engine source listing", level=1)
    para(
        doc,
        "Verbatim listing of unique engine source for patent counsel. Formatting is for identification, not execution. "
        "Website shell source is omitted.",
    )
    for rel, lines, digest, text in stats:
        add_source(doc, rel, lines, digest, text)

    doc.add_heading("8. Counsel checklist", level=1)
    bullets(
        doc,
        [
            "Confirm claimant legal name, inventors, and assignment chain.",
            "Patent: consider a provisional on U-01 through U-16 as a combined takeoff system.",
            "Copyright of the whole website: use a separate full source deposit, not this pack alone.",
            "Trade secret: keep this .docx and the private repository off public GitHub if unpublished.",
            "Dependencies (FastAPI, Vue, Autodesk, OpenAI) are not AutoVAD inventions.",
        ],
    )
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = end.add_run("END OF CONFIDENTIAL PATENT PACK")
    set_run_font(run, size=11, bold=True, color=RED)

    try:
        doc.save(OUT_PATH)
        return OUT_PATH
    except PermissionError:
        alt = OUT_PATH.with_name(OUT_PATH.stem + "_updated.docx")
        doc.save(alt)
        return alt


if __name__ == "__main__":
    print(build())
